"""Blibli spreadsheet style for the Excel exports: Segoe UI, a title block, a white-on-blue header row, banded rows,
horizontal borders only, a frozen header with filters, rupiah without decimals and dates as DD-MM-YYYY.

Every value from the database is written with put(), which keeps text as text: a product name that starts with
"=" can never run as a formula in Excel, and characters Excel cannot store are left out.
"""

from copy import copy
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal

from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.comments import Comment, shape_writer
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.xml.functions import Element


def _note_beside_cell(row, column, height, width, _draw=shape_writer._shape_factory):
    """openpyxl draws every note at the same place, the top left of the sheet. Excel places a note by its anchor
    instead: so put each one beside its own cell, as Excel does for a note made by hand. The anchor counts rows and
    columns from 0 with offsets in pixels; the note's size becomes columns of about 100 px and rows of 20 px."""
    shape = _draw(row, column, height, width)
    anchor = Element(f"{{{shape_writer.excelns}}}Anchor")
    left, top = column + 1, max(row - 1, 0)
    columns, right_offset = divmod(int(15 + width), 100)
    rows, bottom_offset = divmod(int(10 + height), 20)
    anchor.text = f"{left}, 15, {top}, 10, {left + columns}, {right_offset}, {top + rows}, {bottom_offset}"
    shape.find(f"{{{shape_writer.excelns}}}ClientData").insert(2, anchor)  # after MoveWithCells, SizeWithCells
    return shape


if shape_writer._shape_factory.__name__ != _note_beside_cell.__name__:  # once, also if this module is reloaded
    shape_writer._shape_factory = _note_beside_cell

FONT_NAME = "Segoe UI"

# Blibli palette (hex without #)
BLUE = "0072FF"
BAND = "E1F1FF"
TEXT = "28323B"
MUTED = "60686E"
BORDER = "E1E3E4"
TOTAL_FILL = "F1F2F2"
TOTAL_BORDER = "ACB0B4"
POSITIVE = "1F912E"
NEGATIVE = "E93C3C"
WARNING = "BC6E06"
DIFFERS_FILL = "FFF4D1"  # a cell whose value the OCR read differently (amber, as on the website)

FMT = {
    "idr": '"Rp "#,##0;[Red]-"Rp "#,##0',
    "idr_or_dash": '"Rp "#,##0;-"Rp "#,##0;"–"',  # zero as a dash, for amounts that are usually zero
    "number": "#,##0.0##",
    "count": "#,##0",
    "date": "dd-mm-yyyy",
    "datetime": "dd-mm-yyyy hh:mm",
    "month": "mmmm yyyy",
    "month_id": "[$-421]mmmm yyyy",  # the month's name in Indonesian ("Oktober 2026"), whatever Excel's language
}


@dataclass
class Column:
    header: str
    width: float
    fmt: str | None = None
    align: str | None = None  # "right" for numbers, "center" for short codes; text is left aligned
    wrap: bool = False
    group: str | None = None  # a label over neighbouring columns, e.g. "Product" over "Submitted" and "OCR"


@dataclass
class Styled:
    """A cell value with its own look (a status colour, a NULL from the OCR, a cell that differs), optionally a note
    (Excel shows it when pointing at the cell; a red corner marks the cell)."""

    value: object
    color: str | None = None
    fill: str | None = None
    bold: bool = False
    italic: bool = False
    fmt: str | None = None
    link: str | None = None
    note: str | None = None


def put(cell, value):
    """Write a value: text stays text (never a formula), Decimal becomes a number, the rest as it is."""
    if isinstance(value, Decimal):
        value = float(value)
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
        cell.value = value
        cell.data_type = "s"
    else:
        cell.value = value


def link_text(text, linked):
    """The text with only `linked` (e.g. a report code in a sentence) looking like a link: blue and underlined, the
    rest as plain text. For a cell with a link: Excel links the whole cell, so this is only how it looks."""
    before, found, after = (ILLEGAL_CHARACTERS_RE.sub("", part) for part in text.partition(linked))
    # A part of only spaces makes Excel refuse to open the file (see key)
    if not found or (before and not before.strip()) or (after and not after.strip()):
        return text
    plain = InlineFont(rFont=FONT_NAME, sz=11, color=TEXT)
    link = InlineFont(rFont=FONT_NAME, sz=11, color=BLUE, u="single")
    return CellRichText([block for block in (TextBlock(plain, before) if before else None, TextBlock(link, found),
                                             TextBlock(plain, after) if after else None) if block])


def title_block(ws, title, subtitle, column=1):
    """The title and a grey line under it. With frozen columns, start right of them (column=2 when column A is
    frozen): the title then scrolls with the table instead of being cut off at the frozen edge."""
    put(ws.cell(row=1, column=column), title)
    ws.cell(row=1, column=column).font = Font(name=FONT_NAME, size=20, bold=True, color=TEXT)
    put(ws.cell(row=2, column=column), subtitle)
    ws.cell(row=2, column=column).font = Font(name=FONT_NAME, size=9, color=MUTED)
    ws.row_dimensions[1].height = 30


def key(ws, row, column, parts):
    """A small grey line that says what the colours mean, in one cell; each part (Styled, our own text) can have
    the colour it explains, e.g. [Styled("Orange", color=WARNING), Styled(" OCR unsure")]. No part may be only
    spaces: Excel refuses to open the file then."""
    cell = ws.cell(row=row, column=column)
    cell.value = CellRichText([
        TextBlock(
            InlineFont(rFont=FONT_NAME, sz=9, b=part.bold or None, i=part.italic or None, color=part.color or MUTED),
            part.value,
        )
        for part in parts
    ])
    cell.font = Font(name=FONT_NAME, size=9, color=MUTED)


def _align(column, vertical, wrap=False):
    """The column's alignment with a little space on the side the text is against, so neighbours do not touch."""
    horizontal = column.align or "left"
    return Alignment(horizontal=horizontal, vertical=vertical, wrap_text=wrap, indent=1 if horizontal != "center" else 0)


def table(ws, columns, rows, header_row, empty_text, *, bands=None, totals=None, frozen_columns=1):
    """The header, the rows (plain values or Styled) and an optional totals row ({column number: label or formula}).

    bands: per row, True for the light blue band (default: every second row). Columns with a group get its label in
    the row above the header, centred over the neighbours of the same group (centerContinuous instead of merged
    cells, so filters and sorting keep working). The header and the first frozen_columns columns stay in view.
    Returns the number of the last table row (the totals row not counted)."""
    line = Side(style="thin", color=BORDER)
    gap = Side(style="medium", color="FFFFFF")

    if any(column.group for column in columns):
        previous = None
        for index, column in enumerate(columns, start=1):
            if column.group:
                cell = ws.cell(row=header_row - 1, column=index)
                if column.group != previous:
                    put(cell, column.group)
                cell.font = Font(name=FONT_NAME, size=10, bold=True, color=BLUE)
                cell.fill = PatternFill("solid", fgColor=BAND)
                cell.alignment = Alignment(horizontal="centerContinuous", vertical="center")
                cell.border = Border(left=gap if column.group != previous else None)
            previous = column.group
        ws.row_dimensions[header_row - 1].height = 20

    widths = _fitting_widths(columns, rows, totals)
    for index, column in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(index)].width = widths[index - 1]
        cell = ws.cell(row=header_row, column=index)
        put(cell, column.header)
        cell.font = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=BLUE)
        # Top aligned: the filter button sits in the bottom right corner and would cover the end of the text.
        cell.alignment = _align(column, "top", wrap=True)
    ws.row_dimensions[header_row].height = 32

    if not rows:
        cell = ws.cell(row=header_row + 1, column=1)
        put(cell, empty_text)
        cell.font = Font(name=FONT_NAME, size=11, italic=True, color=MUTED)
        last_row = header_row + 1
    else:
        # The look of each kind of cell (column, colours, format, band), made once and copied to the other cells of
        # that kind: making it for every cell took most of the time of a big export (~45 of ~55 s for 5,000 reports).
        looks = {}
        for offset, values in enumerate(rows, start=1):
            row = header_row + offset
            banded = bands[offset - 1] if bands is not None else offset % 2 == 0
            for index, (column, value) in enumerate(zip(columns, values), start=1):
                cell = ws.cell(row=row, column=index)
                styled = value if isinstance(value, Styled) else Styled(value)
                put(cell, styled.value)
                if styled.link:
                    cell.hyperlink = styled.link
                if styled.note:
                    cell.comment = _note(styled.note)
                fmt = styled.fmt or column.fmt
                fmt = fmt if fmt and not isinstance(styled.value, str) else None
                look = (index, styled.bold, styled.italic, styled.color, bool(styled.link), fmt, styled.fill, banded)
                if look in looks:
                    cell._style = copy(looks[look])
                    continue
                cell.font = Font(
                    name=FONT_NAME,
                    size=11,
                    bold=styled.bold,
                    italic=styled.italic,
                    color=styled.color or (BLUE if styled.link else TEXT),
                    underline="single" if styled.link else None,
                )
                if fmt:
                    cell.number_format = fmt
                cell.alignment = _align(column, "top", wrap=column.wrap)
                cell.border = Border(bottom=line)
                if styled.fill or banded:
                    cell.fill = PatternFill("solid", fgColor=styled.fill or BAND)
                # openpyxl keeps a cell's look as one record of numbers (cell._style); copying it is how openpyxl
                # copies a look itself. Private: openpyxl is held below 3.2 in requirements.txt.
                looks[look] = copy(cell._style)
        last_row = header_row + len(rows)

    if totals and rows:
        top = Side(style="thin", color=TOTAL_BORDER)
        for index, column in enumerate(columns, start=1):
            cell = ws.cell(row=last_row + 1, column=index)
            if index in totals:
                cell.value = totals[index]  # our own label or formula, never data from the database
            cell.font = Font(name=FONT_NAME, size=11, bold=True, color=TEXT)
            cell.fill = PatternFill("solid", fgColor=TOTAL_FILL)
            cell.border = Border(top=top)
            cell.alignment = _align(column, "center")
            if column.fmt:
                cell.number_format = column.fmt

    last_column = get_column_letter(len(columns))
    ws.freeze_panes = ws.cell(row=header_row + 1, column=frozen_columns + 1)
    ws.auto_filter.ref = f"A{header_row}:{last_column}{last_row}"
    ws.sheet_view.showGridLines = False
    # Printing: landscape, as wide as one page, the header on every page.
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{header_row - 1 if any(column.group for column in columns) else header_row}:{header_row}"
    return last_row


MAX_TEXT_WIDTH = 60  # a longer text in a column without wrapping runs into its neighbour instead of widening it more


def _fitting_widths(columns, rows, totals):
    """Each column's width: its own, or wider so every value fits (Excel shows #### for a number or date that doesn't,
    and cuts off text). Text in a wrapping column (and the header, which wraps) only needs its longest word to fit, so
    no word is broken; a total (SUBTOTAL) counts as the sum of its column."""
    # The header's filter button takes about 2 more
    widths = [max(column.width, _width(_longest_word(column.header), bold=True) + 2) for column in columns]
    sums = [0.0] * len(columns)
    for values in rows:
        for index, (column, value) in enumerate(zip(columns, values)):
            styled = value if isinstance(value, Styled) else Styled(value)
            if styled.value is None:
                continue
            if isinstance(styled.value, (str, CellRichText)):
                text = str(styled.value)
                characters = _longest_word(text) if column.wrap else _shown_length(text, None)
                widths[index] = max(widths[index], min(_width(characters, styled.bold), MAX_TEXT_WIDTH))
                continue
            if isinstance(styled.value, (int, float, Decimal)) and not isinstance(styled.value, bool):
                sums[index] += abs(float(styled.value))
            widths[index] = max(widths[index], _width(_shown_length(styled.value, styled.fmt or column.fmt), styled.bold))
    for index, total in (totals or {}).items():
        if isinstance(total, str) and total.startswith("="):
            column = columns[index - 1]
            widths[index - 1] = max(widths[index - 1], _width(_shown_length(sums[index - 1], column.fmt), bold=True))
    return widths


def _longest_word(text):
    return max((len(word) for word in text.split()), default=0)


def _width(characters, bold=False):
    """The column width (Excel counts it in Calibri digits) that holds this many characters of Segoe UI 11, a little
    wider than Calibri, with the cell's indent on one side and a margin."""
    return round(characters * (1.3 if bold else 1.2) + 3, 1)


def _shown_length(value, fmt):
    """About how many characters a value takes as Excel shows it with its number format ("Rp 1.358.170",
    "27-10-2026 09:57", "September 2026")."""
    if isinstance(value, str):
        return max((len(line) for line in value.split("\n")), default=0)
    if isinstance(value, (datetime, date)):
        return 14 if fmt and "mmmm" in fmt else len(fmt or "dd-mm-yyyy")
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        number = abs(float(value))
        shown = f"{number:,.3f}".rstrip("0").rstrip(".") if fmt and ".0" in fmt else f"{number:,.0f}"
        return len(shown) + (1 if value < 0 else 0) + (3 if fmt and '"Rp' in fmt else 0)
    return len(str(value))


def _note(text):
    """An Excel note without an author name, sized to its text in pixels (Excel writes notes in Tahoma 9)."""
    lines = text.split("\n")
    note = Comment(text, "")
    note.width = max(len(line) for line in lines) * 6.5 + 20
    note.height = len(lines) * 15 + 14
    return note


def column_range(column_index, first_row, last_row):
    letter = get_column_letter(column_index)
    return f"{letter}{first_row}:{letter}{last_row}"


def excel_month(value, local_zone):
    """The month of a timestamp from the database in the reviewer's time zone, as its first day (a date Excel sorts
    and filters by month; shown with FMT["month"])."""
    moment = excel_time(value, local_zone)
    return moment.date().replace(day=1) if isinstance(moment, datetime) else moment


def excel_time(value, local_zone):
    """A timestamp from the database in the reviewer's time zone, without the zone (Excel has none)."""
    if isinstance(value, datetime):
        if value.tzinfo is not None:
            value = value.astimezone(local_zone)
        return value.replace(tzinfo=None)
    return value
